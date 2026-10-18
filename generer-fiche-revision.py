# -*- coding: utf-8 -*-
"""
Générateur de la fiche de révision du Chapitre 4 · Formulaires HTML5
Format de sortie : .docx (modifiable dans Word / LibreOffice / Google Docs)
Auteur : KELLOUCH Mahdi · ISMO OFPPT
Module : M104 · Développement Digital
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ============================================================
# PALETTE DE COULEURS (charte ISMO · Chapitre 4)
# ============================================================
COLOR_PRIMARY   = RGBColor(0x1E, 0x3A, 0x5F)  # Bleu foncé
COLOR_SECONDARY = RGBColor(0x2C, 0x52, 0x82)  # Bleu moyen
COLOR_ACCENT    = RGBColor(0xD9, 0x77, 0x06)  # Orange
COLOR_LIGHT_BG  = "EEF4FA"                    # Bleu très clair (fond)
COLOR_LIGHT_BG2 = "DBE7F3"                    # Bleu clair (bordures)
COLOR_TEXT      = RGBColor(0x1A, 0x2D, 0x45)  # Texte principal
COLOR_MUTED     = RGBColor(0x4A, 0x6A, 0x8F)  # Texte secondaire
COLOR_CODE_BG   = "F6F8FA"                    # Fond code
COLOR_CODE_TEXT = RGBColor(0x24, 0x29, 0x2E)  # Texte code
COLOR_WARN_BG   = "FFF8E1"                    # Jaune clair
COLOR_OK_BG     = "E8F5EC"                    # Vert clair

# ============================================================
# UTILITAIRES DE MISE EN FORME
# ============================================================
def set_cell_background(cell, color_hex):
    """Applique une couleur de fond à une cellule de tableau."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tc_pr.append(shd)


def set_cell_borders(cell, color_hex="B8CFE5", size="8"):
    """Ajoute une bordure fine à une cellule."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        border = OxmlElement(f'w:{edge}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), size)
        border.set(qn('w:color'), color_hex)
        tc_borders.append(border)
    tc_pr.append(tc_borders)


def add_horizontal_line(paragraph, color="B8CFE5", size="12"):
    """Trace une ligne horizontale sous un paragraphe."""
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), size)
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_paragraph_with_shading(doc, text, bg_color, text_color=COLOR_TEXT,
                                bold=False, italic=False, font_size=10,
                                left_indent=0.3):
    """Ajoute un paragraphe avec un fond coloré (simule un encadré)."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(left_indent)
    p.paragraph_format.right_indent = Cm(left_indent)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)

    # Fond du paragraphe
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), bg_color)
    p_pr.append(shd)

    run = p.add_run(text)
    run.font.size = Pt(font_size)
    run.font.color.rgb = text_color
    run.bold = bold
    run.italic = italic
    return p


def add_code_block(doc, code_text):
    """Insère un bloc de code monospace sur fond clair."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.right_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)

    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), COLOR_CODE_BG)
    p_pr.append(shd)

    # Bordure gauche orange
    p_bdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '24')
    left.set(qn('w:space'), '8')
    left.set(qn('w:color'), 'D97706')
    p_bdr.append(left)
    p_pr.append(p_bdr)

    run = p.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    run.font.color.rgb = COLOR_CODE_TEXT
    return p


def add_section_title(doc, number, title, emoji=""):
    """Ajoute un titre de section avec numéro, emoji et ligne de séparation."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(f"{emoji} {number}. {title}")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = COLOR_PRIMARY
    add_horizontal_line(p, color="B8CFE5", size="12")
    return p


def add_subtitle(doc, text):
    """Sous-titre de niveau 3."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11.5)
    run.font.color.rgb = COLOR_SECONDARY
    return p


def add_bullet(doc, text, level=0):
    """Puce simple."""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Cm(0.7 + level * 0.5)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.size = Pt(10.5)
    run.font.color.rgb = COLOR_TEXT
    return p


def add_info_box(doc, title, content, bg=COLOR_LIGHT_BG,
                 border_color="2C5282"):
    """Encadré pédagogique (info, warn, success)."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.right_indent = Cm(0.3)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)

    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), bg)
    p_pr.append(shd)

    p_bdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '24')
    left.set(qn('w:space'), '8')
    left.set(qn('w:color'), border_color)
    p_bdr.append(left)
    p_pr.append(p_bdr)

    run_title = p.add_run(f"{title} ")
    run_title.bold = True
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = COLOR_PRIMARY

    run_content = p.add_run(content)
    run_content.font.size = Pt(10)
    run_content.font.color.rgb = COLOR_TEXT
    return p


# ============================================================
# GÉNÉRATION DE LA FICHE
# ============================================================
def generate_fiche(output_path="chapitre-4-fiche-revision.docx"):
    doc = Document()

    # ----- Configuration de la page -----
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)

    # ----- Police par défaut -----
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10.5)
    style.font.color.rgb = COLOR_TEXT

    # ============================================================
    # EN-TÊTE
    # ============================================================
    p_badge = doc.add_paragraph()
    p_badge.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p_badge.add_run("  FICHE DE RÉVISION · CHAPITRE 4  ")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = COLOR_PRIMARY
    # Badge avec fond
    p_pr = p_badge._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), COLOR_LIGHT_BG2)
    p_pr.append(shd)

    # Titre principal
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p_title.add_run("Formulaires HTML5")
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = COLOR_PRIMARY

    # Sous-titre
    p_sub = doc.add_paragraph()
    run = p_sub.add_run("Collecter, valider et traiter les données utilisateur")
    run.italic = True
    run.font.size = Pt(12)
    run.font.color.rgb = COLOR_MUTED
    add_horizontal_line(p_sub, color="D97706", size="18")

    # Métadonnées
    p_meta = doc.add_paragraph()
    run = p_meta.add_run("Module M104 · Développement Digital  |  ISMO OFPPT  |  Formateur : KELLOUCH Mahdi  |  2025 – 2026")
    run.font.size = Pt(9)
    run.font.color.rgb = COLOR_MUTED
    p_meta.paragraph_format.space_after = Pt(10)

    # ============================================================
    # 1. ANATOMIE D'UN FORMULAIRE
    # ============================================================
    add_section_title(doc, "1", "Anatomie d'un formulaire", "🧩")

    add_subtitle(doc, "Balise <form> et attributs essentiels")
    add_bullet(doc, "action : URL de destination des données (ex : traitement.php).")
    add_bullet(doc, "method : get (données dans l'URL) ou post (données dans le corps — recommandé).")
    add_bullet(doc, "enctype : type d'encodage. multipart/form-data est obligatoire pour l'envoi de fichiers.")
    add_bullet(doc, "name : nom du formulaire (côté serveur).")
    add_bullet(doc, "target : fenêtre cible (_blank, _self…).")
    add_bullet(doc, "novalidate : désactive la validation HTML5 native.")

    add_info_box(doc, "📌 GET vs POST :",
                 "GET → données visibles dans l'URL, limitées en taille (recherches). "
                 "POST → données invisibles, plus sécurisées (inscriptions).",
                 bg=COLOR_LIGHT_BG, border_color="2C5282")

    add_code_block(doc,
"""<form action="traitement.php" method="post" enctype="multipart/form-data">
  <!-- Champs -->
</form>""")

    # ============================================================
    # 2. CHAMPS DE SAISIE ESSENTIELS
    # ============================================================
    add_section_title(doc, "2", "Champs de saisie essentiels", "⌨️")

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    headers = ["Type", "Usage", "Exemple"]
    for i, h in enumerate(headers):
        set_cell_background(hdr[i], "1E3A5F")
        set_cell_borders(hdr[i], color_hex="1E3A5F")
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    rows_data = [
        ("text", "Texte court", '<input type="text">'),
        ("email", "Adresse email (validation auto)", '<input type="email">'),
        ("password", "Mot de passe masqué", '<input type="password">'),
        ("number", "Nombre (min, max, step)", '<input type="number">'),
        ("date", "Sélecteur de date", '<input type="date">'),
        ("tel", "Téléphone (clavier numérique mobile)", '<input type="tel">'),
        ("url", "Adresse web", '<input type="url">'),
        ("color", "Sélecteur de couleur", '<input type="color">'),
        ("range", "Curseur de valeur", '<input type="range">'),
        ("file", "Upload de fichier (accept)", '<input type="file">'),
        ("search", "Champ de recherche", '<input type="search">'),
        ("time", "Heure", '<input type="time">'),
        ("datetime-local", "Date et heure locales", '<input type="datetime-local">'),
        ("month", "Mois de l'année", '<input type="month">'),
        ("week", "Semaine de l'année", '<input type="week">'),
    ]
    for i, (t, usage, ex) in enumerate(rows_data):
        cells = table.add_row().cells
        if i % 2 == 0:
            for c in cells:
                set_cell_background(c, "F6FAFF")
        for c in cells:
            set_cell_borders(c)
        r0 = cells[0].paragraphs[0].add_run(t)
        r0.font.name = 'Consolas'
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = COLOR_PRIMARY
        r0.bold = True
        r1 = cells[1].paragraphs[0].add_run(usage)
        r1.font.size = Pt(9.5)
        r2 = cells[2].paragraphs[0].add_run(ex)
        r2.font.name = 'Consolas'
        r2.font.size = Pt(9)
        r2.font.color.rgb = COLOR_CODE_TEXT

    add_info_box(doc, "💡 Astuce :",
                 "Le type d'input déclenche le bon clavier sur mobile (email → @, tel → pavé numérique…).",
                 bg=COLOR_LIGHT_BG, border_color="2C5282")

    # ============================================================
    # 3. LISTES, BOUTONS, ZONES DE TEXTE
    # ============================================================
    add_section_title(doc, "3", "Listes, boutons et zones de texte", "🎛️")

    add_subtitle(doc, "Liste déroulante : <select> et <option>")
    add_code_block(doc,
"""<label for="pays">Pays :</label>
<select id="pays" name="pays">
  <option value="">-- Choisir --</option>
  <option value="ma">Maroc</option>
</select>""")

    add_subtitle(doc, "Boutons radio et cases à cocher")
    add_bullet(doc, "Radio : un seul choix possible. Même attribut name pour grouper.")
    add_bullet(doc, "Checkbox : plusieurs choix possibles.")

    add_code_block(doc,
"""<!-- Radio : un seul choix -->
<label><input type="radio" name="genre" value="f"> Femme</label>
<label><input type="radio" name="genre" value="h"> Homme</label>

<!-- Checkbox : plusieurs choix -->
<label><input type="checkbox" name="html" value="html"> HTML5</label>
<label><input type="checkbox" name="css" value="css"> CSS3</label>""")

    add_subtitle(doc, "Zone de texte : <textarea>")
    add_code_block(doc,
"""<label for="message">Message :</label>
<textarea id="message" name="message" rows="4" cols="30"></textarea>""")

    add_subtitle(doc, "Boutons")
    add_code_block(doc,
"""<input type="submit" value="Envoyer">
<input type="reset" value="Réinitialiser">
<button type="button">Annuler</button>""")

    # ============================================================
    # 4. VALIDATION HTML5
    # ============================================================
    add_section_title(doc, "4", "Validation HTML5", "✅")

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    headers = ["Attribut", "Rôle", "Exemple"]
    for i, h in enumerate(headers):
        set_cell_background(hdr[i], "1E3A5F")
        set_cell_borders(hdr[i], color_hex="1E3A5F")
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    valid_data = [
        ("required", "Champ obligatoire", '<input required>'),
        ("pattern", "Expression régulière (regex)", 'pattern="[0-9]{5}"'),
        ("min / max", "Bornes numériques ou dates", 'min="0" max="100"'),
        ("minlength / maxlength", "Longueur du texte", 'minlength="3"'),
        ("step", "Pas d'incrémentation", 'step="0.01"'),
        ("placeholder", "Texte indicatif (disparaît à la saisie)", 'placeholder="Votre nom"'),
        ("autocomplete", "Suggestions du navigateur", 'autocomplete="email"'),
        ("autofocus", "Focus automatique au chargement", '<input autofocus>'),
        ("disabled", "Champ désactivé (NON envoyé)", '<input disabled>'),
        ("readonly", "Lecture seule (envoyé)", '<input readonly>'),
    ]
    for i, (a, role, ex) in enumerate(valid_data):
        cells = table.add_row().cells
        if i % 2 == 0:
            for c in cells:
                set_cell_background(c, "F6FAFF")
        for c in cells:
            set_cell_borders(c)
        r0 = cells[0].paragraphs[0].add_run(a)
        r0.font.name = 'Consolas'
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = COLOR_PRIMARY
        r0.bold = True
        r1 = cells[1].paragraphs[0].add_run(role)
        r1.font.size = Pt(9.5)
        r2 = cells[2].paragraphs[0].add_run(ex)
        r2.font.name = 'Consolas'
        r2.font.size = Pt(9)
        r2.font.color.rgb = COLOR_CODE_TEXT

    add_info_box(doc, "⚠️ Attention :",
                 "disabled empêche l'envoi de la donnée, readonly l'autorise. "
                 "Ces deux attributs sont souvent confondus en EFM.",
                 bg=COLOR_WARN_BG, border_color="EAB308")

    # ============================================================
    # 5. STRUCTURE ET ACCESSIBILITÉ
    # ============================================================
    add_section_title(doc, "5", "Structure et accessibilité", "♿")

    add_subtitle(doc, "Grouper les champs : <fieldset> et <legend>")
    add_code_block(doc,
"""<fieldset>
  <legend>Informations personnelles</legend>
  <label for="nom">Nom :</label>
  <input type="text" id="nom" name="nom">
</fieldset>""")

    add_subtitle(doc, "Associer un label à un champ")
    add_bullet(doc, "Méthode 1 (recommandée) : for sur le <label> + id sur le champ.")
    add_bullet(doc, "Méthode 2 : imbriquer le champ dans le <label>.")

    add_info_box(doc, "✅ Règle W3C :",
                 "Chaque champ doit être associé à un <label>. C'est obligatoire pour "
                 "l'accessibilité (lecteurs d'écran) et améliore l'ergonomie (clic sur le label = focus).",
                 bg=COLOR_OK_BG, border_color="16A34A")

    # ============================================================
    # 6. BONNES PRATIQUES & PIÈGES
    # ============================================================
    add_section_title(doc, "6", "Bonnes pratiques & pièges à éviter", "🎯")

    add_subtitle(doc, "Les 5 règles d'or")
    add_bullet(doc, "Toujours associer un <label> à chaque champ.")
    add_bullet(doc, "Utiliser le bon type d'input (email, tel, date…).")
    add_bullet(doc, "Grouper les champs apparentés avec <fieldset>.")
    add_bullet(doc, "Préférer method=\"post\" pour les données sensibles.")
    add_bullet(doc, "Valider côté client (HTML5) ET côté serveur (PHP, Node…).")

    add_subtitle(doc, "Pièges fréquents en EFM")
    add_info_box(doc, "⚠️ Piège 1 :",
                 "placeholder utilisé comme label → le texte disparaît à la saisie. "
                 "Toujours doubler avec un <label>.",
                 bg=COLOR_WARN_BG, border_color="EAB308")
    add_info_box(doc, "⚠️ Piège 2 :",
                 "Oublier enctype=\"multipart/form-data\" pour l'upload de fichiers → "
                 "seul le nom du fichier est envoyé.",
                 bg=COLOR_WARN_BG, border_color="EAB308")
    add_info_box(doc, "⚠️ Piège 3 :",
                 "Confondre id (pour le label) et name (pour le serveur).",
                 bg=COLOR_WARN_BG, border_color="EAB308")
    add_info_box(doc, "⚠️ Piège 4 :",
                 "Boutons radio sans name commun → on peut en cocher plusieurs.",
                 bg=COLOR_WARN_BG, border_color="EAB308")

    # ============================================================
    # 7. MÉMO EXPRESS · L'ESSENTIEL EN 1 PAGE
    # ============================================================
    add_section_title(doc, "7", "Mémo express · L'essentiel à retenir", "🧠")

    add_info_box(doc, "🔑 À retenir absolument :",
                 " 1) <form action=\"...\" method=\"post\">  |  "
                 "2) <label for=\"x\"> + <input id=\"x\">  |  "
                 "3) required + pattern pour valider  |  "
                 "4) enctype pour les fichiers  |  "
                 "5) fieldset + legend pour grouper.",
                 bg=COLOR_LIGHT_BG2, border_color="1E3A5F")

    # ============================================================
    # PIED DE PAGE
    # ============================================================
    p_footer = doc.add_paragraph()
    p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_footer.paragraph_format.space_before = Pt(18)
    add_horizontal_line(p_footer, color="B8CFE5", size="8")
    run = p_footer.add_run(
        "\nOFPPT · ISMO  |  Module M104 · Développement Digital  |  "
        "Fiche de révision · Chapitre 4  |  KELLOUCH Mahdi · 2025 – 2026"
    )
    run.font.size = Pt(8.5)
    run.font.color.rgb = COLOR_MUTED
    run.italic = True

    # ----- Sauvegarde -----
    doc.save(output_path)
    print(f"✅ Fiche de révision générée : {output_path}")


# ============================================================
# POINT D'ENTRÉE
# ============================================================
if __name__ == "__main__":
    generate_fiche("chapitre-4-fiche-revision.docx")